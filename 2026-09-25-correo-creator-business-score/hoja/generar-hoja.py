"""Genera la plantilla de la hoja de Google (se sube a Drive y se convierte).

    python3 hoja/generar-hoja.py

Pestañas:
  Leads            las columnas que lee y escribe el Apps Script
  Formulario Meta  las preguntas y opciones EXACTAS para el formulario instantáneo
  Cómo funciona    los estados de cada fila y cómo reenviar
Los textos de las preguntas salen de Codigo.gs, para que no se desincronicen.
"""
import json, os, re, subprocess
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
GS = os.path.join(AQUI, '..', 'apps-script', 'Codigo.gs')

# Las preguntas y columnas se leen del propio script (vía Node) para no copiarlas a mano.
datos = json.loads(subprocess.check_output(['node', '-e', '''
const { crearEntorno } = require(process.argv[1]);
const { api } = crearEntorno();
console.log(JSON.stringify({ preguntas: api.PREGUNTAS.map(p => ({ columna: p.columna, titulo: p.titulo, opciones: p.opciones.map(o => o.label) })) }));
''', os.path.join(AQUI, '..', 'pruebas', 'entorno.js')]))
codigo = open(GS, encoding='utf-8').read()
encabezados = re.findall(r"\{ k: '[a-z_]+',\s+h: '([^']+)'", codigo)
assert len(encabezados) == 21, encabezados

MORADO = '6B33F4'
wb = Workbook()

# ---------- Leads ----------
ws = wb.active
ws.title = 'Leads'
ws.append(encabezados)
anchos = {'Fecha': 18, 'ID lead': 20, 'Nombre': 22, 'Correo': 28, 'Teléfono': 18, 'Comunidad': 26, 'Situación': 30, 'Producto': 32,
          'Facturación': 20, 'Problema': 34, 'Origen': 9, 'Puntaje': 9, 'Puntaje real': 12, 'Medalla': 10, 'Caso': 7,
          'Diagnóstico': 28, 'Desenlace': 14, 'Estado': 12, 'Enviado': 18, 'Detalle': 50, 'Intentos': 9}
entrada = set(encabezados[:11])
for i, h in enumerate(encabezados, start=1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=MORADO if h in entrada else '27272A')
    c.alignment = Alignment(vertical='center')
    ws.column_dimensions[get_column_letter(i)].width = anchos.get(h, 14)
ws.row_dimensions[1].height = 24
ws.freeze_panes = 'A2'
# ID y teléfono como texto: sin notación científica ni fórmulas.
for h in ('ID lead', 'Teléfono'):
    col = encabezados.index(h) + 1
    for r in range(2, 2001):
        ws.cell(row=r, column=col).number_format = '@'
est = get_column_letter(encabezados.index('Estado') + 1)
rango = f'A2:{get_column_letter(len(encabezados))}2000'
for valor, color in (('enviado', 'ECFDF3'), ('error', 'FEF3F2'), ('incompleto', 'FFF7E6'), ('revisar', 'FFF7E6'), ('omitido', 'F4F4F5'), ('duplicado', 'F4F4F5')):
    ws.conditional_formatting.add(rango, FormulaRule(formula=[f'${est}2="{valor}"'], fill=PatternFill('solid', fgColor=color)))

# ---------- Formulario Meta ----------
fm = wb.create_sheet('Formulario Meta')
fm.column_dimensions['A'].width = 6
fm.column_dimensions['B'].width = 70
fm.column_dimensions['C'].width = 16
fm.append(['', 'Formulario instantáneo · copia estos textos TAL CUAL (preguntas personalizadas de opción múltiple)', ''])
fm['B1'].font = Font(bold=True, size=13, color=MORADO)
fm.append(['', 'Datos de contacto (preguntas predefinidas de Meta): Nombre completo · Correo electrónico · Número de teléfono', ''])
fm.append([])
for n, p in enumerate(datos['preguntas'], start=1):
    fm.append([f'P{n}', p['titulo'], f'→ columna {p["columna"]}'])
    fm.cell(row=fm.max_row, column=2).font = Font(bold=True)
    fm.cell(row=fm.max_row, column=1).font = Font(bold=True, color=MORADO)
    for o in p['opciones']:
        fm.append(['', '    ' + o, ''])
    fm.append([])

# ---------- Cómo funciona ----------
cf = wb.create_sheet('Cómo funciona')
cf.column_dimensions['A'].width = 14
cf.column_dimensions['B'].width = 100
filas = [
    ('Flujo', 'Meta (formulario instantáneo) → Make → webhook del Apps Script → fila en "Leads" → puntaje y diagnóstico → correo con la medalla.'),
    ('', 'Si otra herramienta escribe la fila directamente, el Apps Script la recoge solo cada 5 minutos.'),
    ('', 'Columnas moradas: las rellena el formulario. Columnas oscuras: las escribe el script. No renombres los encabezados.'),
    ('', ''),
    ('Estado', 'Qué significa'),
    ('(vacío)', 'Pendiente: se enviará en la próxima pasada.'),
    ('enviando', 'Enviándose ahora mismo.'),
    ('enviado', 'Correo enviado. La hora está en "Enviado".'),
    ('error', 'Falló el envío (ver Detalle). Se reintenta solo hasta 3 veces.'),
    ('incompleto', 'Alguna respuesta no coincide con las opciones (ver Detalle). Corrígela y vacía Estado.'),
    ('omitido', 'El correo no es válido. Corrígelo y vacía Estado.'),
    ('duplicado', 'Ese ID de lead ya recibió su correo en otra fila.'),
    ('revisar', 'Se cortó mientras se enviaba: puede que haya salido. Revísalo y vacía Estado para reenviar.'),
    ('', ''),
    ('Reenviar', 'Vacía la celda Estado de la fila (o menú Creator Business Score → Reenviar la fila seleccionada).'),
    ('Puntaje', '"Puntaje" es el que ve el lead (con el relleno de la landing). "Puntaje real" es el que decide calificación, caso y UTMs.'),
]
for a, b in filas:
    cf.append([a, b])
for r in (1, 5, 15, 16):
    cf.cell(row=r, column=1).font = Font(bold=True, color=MORADO)
cf.cell(row=5, column=2).font = Font(bold=True)

salida = os.path.join(AQUI, 'creator-business-score-leads.xlsx')
wb.save(salida)
print('ok', salida)
